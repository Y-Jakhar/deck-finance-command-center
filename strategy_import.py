"""Upload adapter for StrategyLab's unchanged canonical OHLCV CSV contract."""

from __future__ import annotations

import csv
from datetime import date, datetime
from hashlib import sha256
import io
from pathlib import Path

from openpyxl import load_workbook


REQUIRED = ("timestamp", "open", "high", "low", "close", "volume")
MAX_UPLOAD = 3_000_000


def read_upload(raw: bytes, filename: str) -> tuple[list[str], list[list[str]]]:
    if not raw or len(raw) > MAX_UPLOAD:
        raise ValueError("Upload must contain 1 to 3,000,000 bytes")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"), newline="")))
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ValueError("CSV must be valid UTF-8 with parseable rows") from exc
    elif suffix == ".xlsx":
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            sheet = workbook.active
            if sheet.max_column > 50:
                raise ValueError("First release supports at most 50 worksheet columns")
            rows = []
            for row in sheet.iter_rows(values_only=True):
                if any(getattr(value, "data_type", None) == "f" or
                       (isinstance(value, str) and value.startswith("=")) for value in row):
                    raise ValueError("Workbook formulas are unsupported; export reviewed values first")
                rows.append([value.isoformat() if isinstance(value, (date, datetime))
                             else "" if value is None else str(value) for value in row])
                if len(rows) > 10_001:
                    raise ValueError("First release supports at most 10,000 observations")
            workbook.close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("XLSX could not be read; use a non-protected values worksheet") from exc
    else:
        raise ValueError("Upload CSV or XLSX only")
    if not rows:
        raise ValueError("File has no header row")
    headers = [value.strip() for value in rows[0]]
    if not headers or any(not value for value in headers) or len(set(headers)) != len(headers):
        raise ValueError("Header names must be nonblank and unique")
    data = [row for row in rows[1:] if any(value.strip() for value in row)]
    if len(data) > 10_000:
        raise ValueError("First release supports at most 10,000 observations")
    if any(len(row) > len(headers) or any(value.strip() for value in row[len(headers):]) for row in data):
        raise ValueError("A row has more cells than the header")
    return headers, data


def preview(raw: bytes, filename: str) -> dict:
    headers, rows = read_upload(raw, filename)
    lower = {column.lower(): column for column in headers}
    return {"filename": Path(filename).name, "original_sha256": sha256(raw).hexdigest(),
            "columns": headers, "row_count": len(rows), "sample": rows[:5],
            "suggested_mapping": {name: lower.get(name, "") for name in REQUIRED}}


def normalize(raw: bytes, filename: str, mapping: dict[str, str]) -> tuple[str, dict]:
    headers, rows = read_upload(raw, filename)
    if set(mapping) != set(REQUIRED) or any(source not in headers for source in mapping.values()):
        raise ValueError("Map exactly one source column to each required OHLCV field")
    if len(set(mapping.values())) != len(REQUIRED):
        raise ValueError("Each OHLCV field needs a different source column")
    indices = [headers.index(mapping[name]) for name in REQUIRED]
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(REQUIRED)
    for row in rows:
        writer.writerow([(row[index] if index < len(row) else "").strip() for index in indices])
    return output.getvalue(), {"original_filename": Path(filename).name,
                               "original_sha256": sha256(raw).hexdigest(),
                               "column_mapping": mapping, "normalized_from": Path(filename).suffix.lower().lstrip(".")}
