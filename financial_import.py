"""Read-only financial-document candidate extraction; every mapping needs analyst review."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

LABELS = {
    "revenue": ("revenue", "total revenue", "net sales", "sales"),
    "cogs": ("cost of revenue", "cost of goods sold", "cost of sales"),
    "gross_profit": ("gross profit",),
    "operating_expenses": ("operating expenses", "total operating expenses"),
    "ebitda": ("ebitda",),
    "da": ("depreciation and amortization", "depreciation & amortization"),
    "ebit": ("operating income", "ebit"),
    "interest": ("interest expense",),
    "taxes": ("income tax expense", "provision for income taxes"),
    "net_income": ("net income", "net earnings"),
    "cash": ("cash and cash equivalents", "cash equivalents", "cash"),
    "receivables": ("accounts receivable", "receivables"),
    "inventory": ("inventory", "inventories"),
    "current_assets": ("total current assets", "current assets"),
    "fixed_assets": ("property plant and equipment", "property, plant and equipment"),
    "total_assets": ("total assets",),
    "debt": ("total debt", "long-term debt", "long term debt"),
    "payables": ("accounts payable", "payables"),
    "current_liabilities": ("total current liabilities", "current liabilities"),
    "total_liabilities": ("total liabilities",),
    "shareholders_equity": ("total stockholders equity", "shareholders equity", "stockholders equity"),
    "cash_from_operations": ("net cash provided by operating activities", "cash flow from operations", "operating cash flow"),
    "capex": ("capital expenditures", "purchases of property and equipment", "capital expenditure"),
    "cash_from_investing": ("net cash used in investing activities", "cash flow from investing"),
    "cash_from_financing": ("net cash from financing activities", "cash flow from financing"),
    "change_nwc": ("change in net working capital", "changes in working capital"),
    "diluted_shares": ("diluted weighted average shares", "diluted shares", "diluted weighted-average shares"),
    "other_income": ("other income expense net", "other income/(expense), net"),
    "other_current_assets": ("other current assets",),
    "other_assets": ("other non-current assets", "other assets"),
    "other_current_liabilities": ("other current liabilities",),
    "other_liabilities": ("other non-current liabilities", "other liabilities"),
    "total_liabilities_equity": ("total liabilities and shareholders equity", "total liabilities and stockholders equity"),
    "finance_lease_liabilities": ("finance lease liabilities", "finance lease liability"),
    "finance_lease_asset_additions": ("finance lease asset additions",),
    "net_change_cash": ("increase decrease in cash", "net change in cash"),
    "dividends": ("payments for dividends", "dividends paid"),
    "repurchases": ("repurchases of common stock", "share repurchases"),
}


def suggested_key(label: str) -> str | None:
    normalized = re.sub(r"[^a-z0-9]+", " ", label.lower()).strip()
    for key, aliases in LABELS.items():
        if normalized in {re.sub(r"[^a-z0-9]+", " ", alias).strip() for alias in aliases}:
            return key
    return None


def candidates(data: bytes, filename: str) -> dict:
    if len(data) > 10_000_000:
        raise ValueError("Each financial file must be 10 MB or smaller")
    suffix = Path(filename).suffix.lower()
    rows: list[dict] = []
    if suffix == ".csv":
        try:
            table = csv.reader(io.StringIO(data.decode("utf-8-sig")))
            for line, cells in enumerate(table, 1):
                if line > 50_000:
                    break
                _append(rows, filename, f"CSV row {line}", cells)
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ValueError(f"CSV could not be read: {exc}") from exc
    elif suffix == ".xlsx":
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            inspected = 0
            for sheet in workbook.worksheets:
                for line, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                    inspected += 1
                    if inspected > 50_000:
                        break
                    _append(rows, filename, f"{sheet.title}!{line}", cells)
                    if len(rows) >= 500:
                        break
                if inspected >= 50_000 or len(rows) >= 500:
                    break
            workbook.close()
        except Exception as exc:
            raise ValueError(f"Excel workbook could not be read: {exc}") from exc
    elif suffix == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(data))
            for page_number, page in enumerate(reader.pages[:50], 1):
                for line_number, line in enumerate((page.extract_text() or "").splitlines()[:10_000], 1):
                    _append(rows, filename, f"PDF page {page_number}, line {line_number}", [line])
                    if len(rows) >= 500:
                        break
                if len(rows) >= 500:
                    break
        except Exception as exc:
            raise ValueError(f"PDF text could not be extracted: {exc}") from exc
    else:
        raise ValueError("Choose a CSV, XLSX or text-readable PDF file")
    return {
        "filename": filename,
        "format": suffix.lstrip("."),
        "candidates": rows[:500],
        "limitations": "Candidate labels and values are unverified. Review source units, fiscal periods, signs, currency, leases and every mapping before saving. PDF extraction may lose table structure; scanned pages need OCR outside this workflow.",
    }


def _append(rows: list[dict], filename: str, reference: str, values) -> None:
    if len(rows) >= 500:
        return
    cells = [str(value).strip() if value is not None else "" for value in values]
    if not any(cells):
        return
    key = suggested_key(cells[0])
    if not key and len(cells) == 1:
        text = cells[0]
        for label_key, aliases in LABELS.items():
            if any(text.lower().startswith(alias) for alias in aliases):
                key = label_key
                break
    if key:
        rows.append({"reference": f"{filename} · {reference}", "label": cells[0][:300], "suggested_key": key, "raw_values": cells[1:8] if len(cells) > 1 else [cells[0]]})
